"""Build the DOCX demo documents. Usage: py docs/demo-kb/source/build_docx.py <out-dir>"""
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

NAVY = RGBColor(0x1F, 0x2A, 0x44)
TEAL = RGBColor(0x1C, 0x9C, 0x8C)
FOOTER = (
    "Synthetic demonstration document. Cloudshelf is a fictional company; "
    "all policies and numbers are invented."
)


def new_doc(title, meta):
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    brand = doc.add_paragraph()
    run = brand.add_run("CLOUDSHELF · CUSTOMER SUPPORT")
    run.bold = True
    run.font.size = Pt(9)
    run.font.color.rgb = TEAL
    heading = doc.add_heading(title, level=0)
    for r in heading.runs:
        r.font.color.rgb = NAVY
    m = doc.add_paragraph(meta)
    m.runs[0].font.size = Pt(9)
    m.runs[0].italic = True
    return doc


def h(doc, text):
    p = doc.add_heading(text, level=1)
    for r in p.runs:
        r.font.color.rgb = NAVY


def para(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    if bold_prefix:
        p.add_run(bold_prefix).bold = True
    p.add_run(text)
    return p


def bullets(doc, items, style="List Bullet"):
    for item in items:
        doc.add_paragraph(item, style=style)


def table(doc, rows):
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    t.style = "Light Grid Accent 1"
    for i, row in enumerate(rows):
        for j, value in enumerate(row):
            cell = t.cell(i, j)
            cell.text = value
            if i == 0:
                for r in cell.paragraphs[0].runs:
                    r.bold = True
    doc.add_paragraph()


def footer(doc):
    p = doc.add_paragraph(FOOTER)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.runs[0].font.size = Pt(8)
    p.runs[0].italic = True


def refund_procedure(out):
    doc = new_doc(
        "Refund Procedure",
        "Procedure CS-PRO-007 · Version 5 · Effective September 1, 2026 · Owner: Billing Operations",
    )
    h(doc, "Purpose")
    para(doc, "Use this procedure when a customer asks for a refund on a paid plan.")
    h(doc, "Eligibility")
    para(doc, "A charge can be refunded by a Tier 1 agent only when all of these are true:")
    bullets(doc, [
        "The customer is within the 30-day refund window, counted from the charge date.",
        "The payment has cleared and is not a pending authorization.",
        "The customer has not already received a refund for the same charge.",
    ])
    para(doc, "Anything else is a courtesy refund. A Tier 2 supervisor must approve any courtesy refund.", "Outside these rules: ")
    h(doc, "Procedure")
    bullets(doc, [
        "Verify the customer's identity using two factors.",
        "Open the most recent charge in the billing console.",
        "Select Issue Refund.",
        "Choose Full or Partial. A partial refund needs the amount in dollars and cents.",
        "Add a one-line reason code (RF-01, RF-02, RF-03 or RF-05).",
        "Submit the refund. It will post to the original card within 5-7 business days.",
    ], style="List Number")
    h(doc, "Reason codes")
    table(doc, [
        ["Code", "Use when", "Approval"],
        ["RF-01", "Service not as described", "None"],
        ["RF-02", "Duplicate charge", "None; also open a Billing Operations ticket"],
        ["RF-03", "Unintended purchase or renewal", "None"],
        ["RF-05", "Courtesy refund", "Tier 2 supervisor approval ID required"],
    ])
    para(doc, "RF-04 is retired. Enterprise outage credits are not refunds: they go on the next invoice through Credits > Outage credit in the billing console, never to a card (see the Outages and Service Credits FAQ).", "Outage credits: ")
    h(doc, "Escalation and approval")
    para(doc, "If the customer disputes the amount or the refund window, escalate the request to a Tier 2 supervisor within 15 minutes.")
    para(doc, "Do not issue a courtesy refund without approval from a Tier 2 supervisor.")
    para(doc, "Refunds go only to the original payment method. Never refund to a different card, a bank transfer or a gift card.")
    footer(doc)
    doc.save(out / "refund-procedure.docx")


def identity_verification(out):
    doc = new_doc(
        "Identity Verification Procedure",
        "Procedure CS-PRO-002 · Version 8 · Effective October 1, 2026 · Owner: Security Response",
    )
    h(doc, "When to verify")
    para(doc, "Verify the caller before you discuss, view aloud or change anything on an account. General questions about plans and prices need no verification.")
    h(doc, "Accepted factors")
    para(doc, "Ask for two different factors from this list:")
    table(doc, [
        ["Factor", "How to check it"],
        ["One-time code", "Send a 6-digit code to the account email from the console; the caller reads it back"],
        ["Account PIN", "6-digit PIN only; 4-digit PINs stopped working on October 1, 2026"],
        ["Billing ZIP code", "Must match the ZIP of the card on file"],
        ["Last 4 digits of the card on file", "Visible in the console; the caller states them"],
    ])
    h(doc, "Never ask for")
    bullets(doc, [
        "The full card number or the card security code (CVV).",
        "The account password.",
        "Social Security numbers or any government ID number.",
    ])
    h(doc, "When verification fails")
    para(doc, "If the caller fails verification twice, stop account-specific discussion and point them to self-service recovery from the sign-in page. Do not hint which factor was wrong.")
    para(doc, "If you suspect an account takeover (for example, the caller wants to change the account email and cannot verify), lock the account and escalate to Security Response within 1 hour.")
    h(doc, "Enterprise accounts")
    para(doc, "Only the account administrator can authorize billing changes, cancellation or deletion for an Enterprise organization. Ordinary team members can be verified for their own user settings only.")
    footer(doc)
    doc.save(out / "identity-verification.docx")


if __name__ == "__main__":
    out = Path(sys.argv[1])
    out.mkdir(parents=True, exist_ok=True)
    refund_procedure(out)
    identity_verification(out)
    print("docx done")
